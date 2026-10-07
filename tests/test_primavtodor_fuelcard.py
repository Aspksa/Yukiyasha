"""Monthly fuel card: one sheet per driver, live formulas, totals consistent with waybills."""

import io
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from tests.test_primavtodor_records import make_driver, make_vehicle, make_waybill
from yukiyasha.config import Settings
from yukiyasha.modules.disk import DiskModule
from yukiyasha.modules.primavtodor import PrimavtodorModule, PrintNotAvailableError
from yukiyasha.modules.primavtodor.fuelcard import FuelCard, fill_fuel_cards
from yukiyasha.web.app import create_app


@pytest.fixture
def module(tmp_path: Path) -> PrimavtodorModule:
    disk = DiskModule(tmp_path / "disk")
    disk.start()
    primavtodor = PrimavtodorModule(disk)
    primavtodor.start()
    return primavtodor


def book(content: bytes) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(content))


def month_with_two_drivers(module: PrimavtodorModule) -> dict:
    vehicle = make_vehicle(module, plate="А123ВС125", model="TOYOTA TEST", norm_summer="10")
    first = make_driver(module, vehicle["id"], full_name="Иванов Иван Иванович",
                        personnel_number="1", fuel_card_number="7001 0000 0001")  # fmt: skip
    second = make_driver(module, vehicle["id"], full_name="Петров Пётр Петрович",
                         personnel_number="2", fuel_card_number="7001000000002")  # fmt: skip
    one = make_waybill(module, first, vehicle, number="1", date="2026-10-02", season="summer",
                       odometer_out=1000, odometer_in=1100, fuel_out=20, fuel_in=15)  # fmt: skip
    two = make_waybill(module, first, vehicle, number="2", date="2026-10-05", season="summer",
                       odometer_out=1100, odometer_in=1300, fuel_out=15, fuel_in=10)  # fmt: skip
    make_waybill(module, second, vehicle, number="3", date="2026-10-09", season="summer",
                 odometer_out=1300, odometer_in=1350, fuel_out=10, fuel_in=8)  # fmt: skip
    module.data.create("fuel", {"waybill_id": one["id"], "date": "2026-10-02", "liters": 5})
    module.data.create("fuel", {"waybill_id": two["id"], "date": "2026-10-05", "liters": 10})
    return vehicle


def test_a_sheet_per_driver_with_the_cards_figures(module: PrimavtodorModule) -> None:
    vehicle = month_with_two_drivers(module)

    content, name = module.fuel_cards(vehicle["id"], "2026-10")

    workbook = book(content)
    assert name == "Карточка ГСМ А123ВС125 2026-10.xlsx"
    assert len(workbook.sheetnames) == 2 and all("Иванов" in n or "Петров" in n
                                                  for n in workbook.sheetnames)  # fmt: skip
    ivanov = workbook.worksheets[0]
    assert ivanov["B2"].value == "TOYOTA TEST А123ВС125"
    assert ivanov["D2"].value == "Октябрь 2026" and ivanov["F2"].value == "Иванов И. И."
    assert ivanov["C4"].value == "700100000001"  # the card, whitespace removed
    assert ivanov["D5"].value == 300  # 100 + 200 km
    assert ivanov["D6"].value == 20  # fuel at the first departure
    assert [ivanov.cell(r, 2).value for r in (9, 10, 11)] == [5, 10, None]  # fill-ups
    # consumption: (20+5-15) + (15+10-10) = 25 l; the opening 20 l is burnt first
    assert ivanov["G8"].value == 25
    assert (ivanov["E9"].value, ivanov["E10"].value) == (20, 5)
    assert ivanov["B33"].value == "Остаток на 31.10.2026"
    assert ivanov["B31"].value == "(10*км/100)" and ivanov["E31"].value == "=10*D5/100"
    assert ivanov["B34"].value == "=D6+B28-E28"  # formulas stay live

    petrov = workbook.worksheets[1]
    assert petrov["D5"].value == 50 and petrov["G8"].value == 2 and petrov["B9"].value is None


def test_closing_balance_equals_the_last_waybills_remainder(module: PrimavtodorModule) -> None:
    vehicle = month_with_two_drivers(module)
    sheet = book(module.fuel_cards(vehicle["id"], "2026-10")[0]).worksheets[0]

    income = sum(sheet.cell(r, 2).value or 0 for r in range(9, 28))
    expense = sum(sheet.cell(r, 5).value or 0 for r in range(9, 28))

    assert sheet["D6"].value + income - expense == 10  # 20 + 15 - 25, the remainder of waybill 2


def test_open_waybills_do_not_count_and_a_month_without_waybills_is_refused(
    module: PrimavtodorModule,
) -> None:
    vehicle = make_vehicle(module)
    driver = make_driver(module, vehicle["id"])
    make_waybill(module, driver, vehicle, number="1", date="2026-10-02")  # still open
    sheet = book(module.fuel_cards(vehicle["id"], "2026-10")[0]).worksheets[0]

    assert sheet["D5"].value == 0 and sheet["G8"].value == 0
    with pytest.raises(PrintNotAvailableError, match="нет путевых листов"):
        module.fuel_cards(vehicle["id"], "2026-09")
    with pytest.raises(PrintNotAvailableError, match="ГГГГ-ММ"):
        module.fuel_cards(vehicle["id"], "октябрь")


def test_more_fill_ups_than_lines_are_summed_into_the_last_line() -> None:
    card = FuelCard(vehicle="Т А1", month=__import__("datetime").date(2026, 10, 1), driver="А Б",
                    card_number="1", distance_km=10, opening=0, fillups=[1.0] * 25,
                    consumption=25.0)  # fmt: skip

    sheet = book(fill_fuel_cards([card])).worksheets[0]

    column = [sheet.cell(r, 2).value for r in range(9, 28)]
    assert column[:18] == [1] * 18 and column[18] == 7  # 18 single lines + the remaining 7
    assert sum(column) == 25


def test_mixed_seasons_use_the_sum_of_the_waybill_norms(module: PrimavtodorModule) -> None:
    vehicle = make_vehicle(module, norm_summer="10", norm_winter="20")
    driver = make_driver(module, vehicle["id"])
    make_waybill(module, driver, vehicle, number="1", date="2026-03-30", season="winter",
                 odometer_out=0, odometer_in=100, fuel_out=30, fuel_in=10)  # fmt: skip
    make_waybill(module, driver, vehicle, number="2", date="2026-03-31", season="summer",
                 odometer_out=100, odometer_in=200, fuel_out=10, fuel_in=0)  # fmt: skip

    sheet = book(module.fuel_cards(vehicle["id"], "2026-03")[0]).worksheets[0]

    assert sheet["B31"].value == "(по путевым листам)" and sheet["E31"].value == 30  # 20 + 10


def test_blank_template_has_no_personal_data() -> None:
    sheet = book(fill_fuel_cards([FuelCard("", __import__("datetime").date(2026, 1, 1), "", "",
                                           0, 0)])).worksheets[0]
    texts = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    assert "Лушкин" not in texts and "5019801" not in texts and "ШМЕЛЬ" not in texts.upper()


@pytest.fixture
def client(tmp_path: Path):
    app = create_app(Settings(disk_dir=tmp_path / "disk"))
    with TestClient(app, base_url="http://127.0.0.1:8000") as test_client:
        yield test_client


def test_fuel_card_over_http(client: TestClient) -> None:
    vehicle = client.post("/api/primavtodor/records/vehicles",
                          json={"plate": "А1", "model": "Т"}).json()  # fmt: skip
    driver = client.post("/api/primavtodor/records/employees",
                         json={"full_name": "Иванов Иван", "is_driver": True}).json()  # fmt: skip
    client.post(
        "/api/primavtodor/records/waybills",
        json={"number": "1", "date": "2026-10-07", "driver_id": driver["id"],
              "vehicle_id": vehicle["id"], "odometer_out": 5},
    )  # fmt: skip
    url = f"/api/primavtodor/vehicles/{vehicle['id']}/fuel-card"

    ok = client.get(url, params={"month": "2026-10"})
    empty = client.get(url, params={"month": "2026-01"})

    assert ok.status_code == 200 and ok.headers["content-type"].startswith("application/vnd")
    assert "filename*=UTF-8''" in ok.headers["content-disposition"]
    assert empty.status_code == 422 and "нет путевых листов" in empty.json()["detail"]["message"]
    assert client.get("/api/primavtodor/vehicles/veh-00000000/fuel-card",
                      params={"month": "2026-10"}).status_code == 404  # fmt: skip
