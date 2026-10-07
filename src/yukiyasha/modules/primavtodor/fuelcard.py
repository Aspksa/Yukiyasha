"""Monthly fuel card of a vehicle («Карточка расхода ГСМ»): one sheet per driver.

The template (``forms/fuel_card.xlsx``) is the organisation's own blank card. It keeps its
formulas, so the totals and the closing balance stay live when someone edits the file:
``Σ расход``, ``расход по норме`` and ``остаток = на начало + Σ приход − Σ расход``.

Приход is the list of fill-ups (litres). Расход follows the card's own convention: the opening
balance is burnt first, the rest of the month's actual consumption goes into the second line,
so the column adds up to the real consumption.
"""

import io
import re
from dataclasses import dataclass, field
from datetime import date
from importlib import resources

import openpyxl

from yukiyasha.modules.primavtodor.errors import PrintNotAvailableError
from yukiyasha.modules.primavtodor.printing import short_name

MONTH_NAMES = (
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
)  # fmt: skip
FIRST_ROW, LAST_ROW = 9, 27
CAPACITY = LAST_ROW - FIRST_ROW + 1
BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


@dataclass(slots=True)
class FuelCard:
    """One driver's card for one vehicle and month. All numbers are litres or kilometres."""

    vehicle: str  # «модель госномер»
    month: date  # any day of the month
    driver: str
    card_number: str
    distance_km: int
    opening: float
    fillups: list[float] = field(default_factory=list)  # in date order
    consumption: float = 0.0
    norm_rate: float | None = None  # l per 100 km when one rate covers the whole month
    norm_total: float | None = None  # sum of the waybills' norms when the rates differ


def _number(value: float) -> float | int:
    rounded = round(value, 3)
    return int(rounded) if rounded == int(rounded) else rounded


def _end_of_month(day: date) -> date:
    first = day.replace(day=1)
    following = first.replace(year=first.year + 1, month=1) if first.month == 12 else first.replace(
        month=first.month + 1
    )
    return date.fromordinal(following.toordinal() - 1)


def _sheet_title(card: FuelCard, taken: set[str]) -> str:
    plate = (card.vehicle.split() or [""])[-1]
    base = BAD_SHEET_CHARS.sub(" ", f"{plate} {short_name(card.driver)}")
    base = " ".join(base.split())[:28] or "Карточка"
    title, suffix = base, 2
    while title in taken:
        title, suffix = f"{base[:26]} {suffix}", suffix + 1
    return title


def fill_fuel_cards(cards: list[FuelCard]) -> bytes:
    """A workbook with one filled card per entry of ``cards`` (at least one)."""
    if not cards:
        raise PrintNotAvailableError("Нет данных для карточки")
    template = resources.files("yukiyasha.modules.primavtodor").joinpath("forms/fuel_card.xlsx")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(template.read_bytes()))
        blank = workbook["Карточка"]
    except (OSError, KeyError, ValueError) as exc:
        raise PrintNotAvailableError("Не удалось открыть бланк карточки ГСМ") from exc

    taken: set[str] = set()
    for card in cards:
        sheet = workbook.copy_worksheet(blank)
        sheet.title = _sheet_title(card, taken)
        taken.add(sheet.title)
        _fill(sheet, card)
    del workbook[blank.title]
    workbook.active = 0
    workbook.properties.title = "Карточка расхода ГСМ"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _fill(sheet, card: FuelCard) -> None:
    month = card.month
    sheet["B2"] = card.vehicle
    sheet["D2"] = f"{MONTH_NAMES[month.month - 1]} {month.year}"
    sheet["F2"] = short_name(card.driver)
    sheet["C4"] = card.card_number or None
    sheet["D5"] = card.distance_km
    sheet["B6"] = f"Остаток на {month.replace(day=1):%d.%m.%Y}"
    sheet["D6"] = _number(card.opening)
    sheet["B33"] = f"Остаток на {_end_of_month(month):%d.%m.%Y}"

    fills = [_number(liters) for liters in card.fillups]
    if len(fills) > CAPACITY:  # the card has 19 lines: the tail is summed into the last one
        fills = [*fills[: CAPACITY - 1], _number(sum(card.fillups[CAPACITY - 1 :]))]
    for offset, liters in enumerate(fills):
        sheet.cell(FIRST_ROW + offset, 2).value = liters

    burnt_from_opening = min(card.opening, max(card.consumption, 0.0))
    sheet.cell(FIRST_ROW, 5).value = _number(burnt_from_opening)
    rest = card.consumption - burnt_from_opening
    if rest > 0:
        sheet.cell(FIRST_ROW + 1, 5).value = _number(rest)
    sheet["G8"] = _number(card.consumption)

    if card.norm_rate is not None:
        rate = _number(card.norm_rate)
        sheet["B31"] = f"({rate}*км/100)"
        sheet["E31"] = f"={rate}*D5/100"
    elif card.norm_total is not None:
        sheet["B31"] = "(по путевым листам)"
        sheet["E31"] = _number(card.norm_total)
