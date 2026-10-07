"""Timesheet as the organisation's own form Т-12 («Табель учета рабочего времени»).

The template (``forms/timesheet_t12.xlsx``) is the blank form with all names and marks removed.
Every employee takes two lines: the marks (Я, РВ, ОТ, ...) and the hours. Weekends and holidays
of the production calendar are painted in the day columns. The per-half-month and monthly
totals keep the form's own formulas, so they stay live when someone edits the file. A form
holds 15 employees; more are added by repeating the last block.
"""

import io
from dataclasses import dataclass, field
from datetime import date
from importlib import resources

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from yukiyasha.modules.primavtodor.errors import PrintNotAvailableError

BLOCKS = 15  # employee blocks (two lines each) on the blank form, rows 16-45
FIRST_ROW = 16
TABLE_END = FIRST_ROW + 2 * BLOCKS  # first line below the table; two spacer lines follow
SIGNATURE_ROW = TABLE_END + 2  # the line «Ответственное лицо»
FOOTER_ROWS = 6  # spacer lines and the signature lines, moved down as a whole
LAST_COLUMN = 256
FIRST_HALF, SECOND_HALF = 33, 100  # first column of day 1 and of day 16; 4 columns per day
FORMULA_COLUMNS = ("CO", "FH", "FO", "FV")
OFF_FILL = PatternFill("solid", fgColor="FFFFFF00")
OFF_WORK_FILL = PatternFill("solid", fgColor="FFFFCC00")
NO_FILL = PatternFill(fill_type=None)
MONTHS_GENITIVE = (
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
)  # fmt: skip


@dataclass(slots=True)
class TimesheetPerson:
    name: str  # «Иванов И. И. Водитель»
    personnel_number: str = ""
    marks: dict[int, str] = field(default_factory=dict)  # day -> code
    hours: dict[int, int] = field(default_factory=dict)  # day -> hours worked


@dataclass(slots=True)
class TimesheetForm:
    month: date
    days_in_month: int
    off_days: set[int]
    org_name: str = ""
    unit: str = ""
    persons: list[TimesheetPerson] = field(default_factory=list)
    composer_title: str = ""
    composer: str = ""
    head_title: str = ""
    head: str = ""
    compiled: date | None = None


def _column(day: int) -> int:
    return FIRST_HALF + 4 * (day - 1) if day <= 15 else SECOND_HALF + 4 * (day - 16)


def _shifted(merged: CellRange, rows: int) -> str:
    moved = CellRange(merged.coord)
    moved.shift(row_shift=rows)  # shifts in place
    return moved.coord


def _copy_block(sheet: Worksheet, source_top: int, target_top: int) -> None:
    """Repeat a two-row employee block (values as formulas, styles, merges, heights)."""
    for offset in (0, 1):
        source, target = source_top + offset, target_top + offset
        sheet.row_dimensions[target].height = sheet.row_dimensions[source].height
        for column in range(1, LAST_COLUMN + 1):
            origin = sheet.cell(source, column)
            cell = sheet.cell(target, column)
            cell._style = origin._style
            if isinstance(origin.value, str) and origin.value.startswith("="):
                cell.value = Translator(origin.value, origin.coordinate).translate_formula(
                    cell.coordinate
                )
    for merged in list(sheet.merged_cells.ranges):
        if merged.min_row == source_top or merged.min_row == source_top + 1:
            sheet.merge_cells(_shifted(merged, target_top - source_top))


def _make_room(sheet: Worksheet, extra: int) -> None:
    """Move the signature lines down and repeat the last employee block ``extra`` times."""
    shift = 2 * extra
    area = range(TABLE_END, TABLE_END + FOOTER_ROWS)
    moving = [m for m in sheet.merged_cells.ranges if m.min_row in area]
    heights = {r: sheet.row_dimensions[r].height for r in area}
    for merged in moving:
        sheet.unmerge_cells(merged.coord)
    sheet.move_range(f"A{area[0]}:IV{area[-1]}", rows=shift)
    for merged in moving:
        sheet.merge_cells(_shifted(merged, shift))
    for row, height in heights.items():
        sheet.row_dimensions[row + shift].height = height
    last_top = TABLE_END - 2
    for index in range(extra):
        _copy_block(sheet, last_top, TABLE_END + 2 * index)


def fill_timesheet(data: TimesheetForm) -> bytes:
    template = resources.files("yukiyasha.modules.primavtodor").joinpath("forms/timesheet_t12.xlsx")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(template.read_bytes()))
        sheet = workbook["Табель"]
    except (OSError, KeyError, ValueError) as exc:
        raise PrintNotAvailableError("Не удалось открыть бланк табеля") from exc

    count = max(len(data.persons), 1)
    extra = max(0, count - BLOCKS)
    if extra:
        _make_room(sheet, extra)
    signature = SIGNATURE_ROW + 2 * extra

    month = data.month
    sheet["I2"] = data.org_name or None
    sheet["I4"] = data.unit or None
    last_day = date(month.year, month.month, data.days_in_month)
    sheet["GF8"], sheet["GY8"] = month, last_day
    compiled = data.compiled
    sheet["FI8"] = compiled
    sheet[f"U{signature}"] = data.composer_title or None
    sheet[f"BW{signature}"] = data.composer or None
    sheet[f"EQ{signature}"] = data.head_title or None
    sheet[f"GJ{signature}"] = data.head or None
    if compiled:
        sheet[f"HH{signature}"] = f"{compiled.day:02d}"
        sheet[f"HM{signature}"] = MONTHS_GENITIVE[compiled.month - 1]
        sheet[f"HX{signature}"] = "20"
        sheet[f"IB{signature}"] = f"{compiled.year % 100:02d}"

    # header line of the day numbers: paint the days off
    for day in range(1, 32):
        in_month = day <= data.days_in_month
        for column in range(_column(day), _column(day) + 4):
            sheet.cell(12, column).fill = OFF_FILL if in_month and day in data.off_days else NO_FILL

    blocks = BLOCKS + extra
    for block in range(blocks):
        top = FIRST_ROW + 2 * block
        person = data.persons[block] if block < len(data.persons) else None
        sheet[f"A{top}"] = block + 1 if person else None
        sheet[f"F{top}"] = person.name if person else None
        sheet[f"V{top}"] = (person.personnel_number or None) if person else None
        if person is None:  # an empty line of the form shows no totals
            for column in FORMULA_COLUMNS:
                sheet[f"{column}{top}"] = None
            sheet[f"CO{top + 1}"] = None  # FO and FV span both lines, CO and FH have one each
            sheet[f"FH{top + 1}"] = None
        for day in range(1, 32):
            code = person.marks.get(day) if person and day <= data.days_in_month else None
            hours = person.hours.get(day) if person and day <= data.days_in_month else None
            if day <= data.days_in_month and day in data.off_days:
                fill = OFF_WORK_FILL if code else OFF_FILL
            else:
                fill = NO_FILL
            for row in (top, top + 1):
                for column in range(_column(day), _column(day) + 4):
                    sheet.cell(row, column).fill = fill
            sheet.cell(top, _column(day)).value = code
            sheet.cell(top + 1, _column(day)).value = hours or None

    # the blank form was painted for one particular month: clear what is left below the table
    for row in range(FIRST_ROW + 2 * blocks, sheet.max_row + 1):
        for day in range(1, 32):
            for column in range(_column(day), _column(day) + 4):
                sheet.cell(row, column).fill = NO_FILL

    workbook.properties.title = f"Табель {month:%m.%Y}"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
