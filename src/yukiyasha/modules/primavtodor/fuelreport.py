"""Monthly «Анализ расхода ГСМ»: one row per vehicle, diesel and petrol in separate columns.

The template (``forms/fuel_report.xlsx``) is the organisation's own blank report. Per vehicle:
mileage, opening balance, fill-ups, actual consumption, consumption by norm and the closing
balance (opening + fill-ups - actual), each in the diesel (ДТ) or the petrol column. The totals
row uses live ``SUM`` formulas. The second balance block («остаток на конец месяца» measured in
the tank) stays empty on purpose: it is filled in by hand.
"""

import io
from copy import copy
from dataclasses import dataclass
from datetime import date
from importlib import resources

import openpyxl

from yukiyasha.modules.primavtodor.errors import PrintNotAvailableError

MONTH_NAMES_LOWER = (
    "январь", "февраль", "март", "апрель", "май", "июнь",
    "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь",
)  # fmt: skip
FIRST_ROW, TEMPLATE_ROWS = 9, 18  # the blank report has 18 vehicle lines (rows 9-26)
COLUMNS = {"opening": "F", "fillups": "H", "actual": "J", "norm": "L", "closing": "N"}
NEXT = {"F": "G", "H": "I", "J": "K", "L": "M", "N": "O"}  # petrol column of each pair
NUMERIC = "EFGHIJKLMNO"


@dataclass(slots=True)
class VehicleRow:
    plate: str
    model: str
    kind: str  # легковой / спец. / грузовой
    drivers: str
    petrol: bool  # False: diesel column
    distance_km: int = 0
    opening: float = 0.0
    fillups: float = 0.0
    actual: float = 0.0
    norm: float = 0.0
    note: str = ""

    @property
    def closing(self) -> float:
        return self.opening + self.fillups - self.actual


@dataclass(slots=True)
class ReportSigners:
    approver_title: str = ""
    approver: str = ""
    composer_title: str = ""
    composer: str = ""


def _number(value: float) -> float | int:
    rounded = round(value, 2)
    return int(rounded) if rounded == int(rounded) else rounded


def _copy_style(source, target) -> None:
    target._style = copy(source._style)


def fill_fuel_report(
    rows: list[VehicleRow], month: date, signers: ReportSigners
) -> bytes:
    if not rows:
        raise PrintNotAvailableError("За этот месяц нет данных по машинам")
    template = resources.files("yukiyasha.modules.primavtodor").joinpath("forms/fuel_report.xlsx")
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(template.read_bytes()))
        sheet = workbook["Анализ"]
    except (OSError, KeyError, ValueError) as exc:
        raise PrintNotAvailableError("Не удалось открыть бланк «Анализ расхода ГСМ»") from exc

    extra = max(0, len(rows) - TEMPLATE_ROWS)
    total_row = FIRST_ROW + TEMPLATE_ROWS  # 27 in the blank
    if extra:
        sheet.unmerge_cells(f"A{total_row}:D{total_row}")
        sheet.move_range(f"A{total_row}:R{total_row + 4}", rows=extra)
        for r in range(total_row + 4 + extra, total_row - 1, -1):
            sheet.row_dimensions[r].height = sheet.row_dimensions[r - extra].height
        for r in range(FIRST_ROW + TEMPLATE_ROWS, FIRST_ROW + TEMPLATE_ROWS + extra):
            sheet.row_dimensions[r].height = sheet.row_dimensions[FIRST_ROW].height
            for column in range(1, 19):
                _copy_style(sheet.cell(FIRST_ROW, column), sheet.cell(r, column))
        total_row += extra
        sheet.merge_cells(f"A{total_row}:D{total_row}")

    sheet["F6"] = f"за {MONTH_NAMES_LOWER[month.month - 1]} {month.year} года"
    sheet["I1"] = f"Согласовано:\n{signers.approver_title}    {signers.approver}".rstrip()
    sheet.cell(total_row + 2, 4).value = signers.composer_title or None
    sheet.cell(total_row + 2, 5).value = signers.composer or None

    for offset, row in enumerate(rows):
        r = FIRST_ROW + offset
        sheet.cell(r, 1).value = row.plate
        sheet.cell(r, 2).value = row.model
        sheet.cell(r, 3).value = row.kind
        sheet.cell(r, 4).value = row.drivers
        sheet.cell(r, 5).value = row.distance_km
        values = {
            "opening": row.opening, "fillups": row.fillups, "actual": row.actual,
            "norm": row.norm, "closing": row.closing,
        }  # fmt: skip
        for key, value in values.items():
            column = COLUMNS[key]
            if row.petrol:
                column = NEXT[column]
            if value or key == "closing":
                sheet[f"{column}{r}"] = _number(value)
        sheet.cell(r, 18).value = row.note or None

    last = FIRST_ROW + len(rows) - 1 if extra else FIRST_ROW + TEMPLATE_ROWS - 1
    for column in NUMERIC:
        sheet[f"{column}{total_row}"] = f"=SUM({column}{FIRST_ROW}:{column}{last})"

    workbook.properties.title = "Анализ расхода ГСМ"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
