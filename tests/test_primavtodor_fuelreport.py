"""Monthly «Анализ расхода ГСМ»: a row per vehicle, diesel/petrol columns, live totals."""

import io
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from tests.test_primavtodor_records import make_driver, make_vehicle, make_waybill
from yukiyasha.config import Settings
from yukiyasha.modules.disk import DiskModule
from yukiyasha.modules.primavtodor import PrimavtodorModule, PrintNotAvailableError
from yukiyasha.modules.primavtodor.fuelreport import ReportSigners, VehicleRow, fill_fuel_report
from yukiyasha.web.app import create_app


@pytest.fixture
def module(tmp_path: Path) -> PrimavtodorModule:
    disk = DiskModule(tmp_path / "disk")
    disk.start()
    primavtodor = PrimavtodorModule(disk)
    primavtodor.start()
    return primavtodor


def sheet_of(content: bytes):
    return openpyxl.load_workbook(io.BytesIO(content))["Анализ"]


def fill(module: PrimavtodorModule) -> None:
    diesel = make_vehicle(module, plate="Д111ДД125", model="КАМАЗ", fuel_type="ДТ",
                          norm_summer="30")  # fmt: skip
    petrol = make_vehicle(module, plate="Б222ББ125", model="TOYOTA", fuel_type="АИ-95",
                          norm_summer="10")  # fmt: skip
    idle = make_vehicle(module, plate="Х000ХХ125", model="Без листов")
    assert idle
    ivan = make_driver(module, diesel["id"], full_name="Иванов Иван Иванович",
                       personnel_number="1", fuel_card_number="1")  # fmt: skip
    petr = make_driver(module, petrol["id"], full_name="Петров Пётр Петрович",
                       personnel_number="2", fuel_card_number="2")  # fmt: skip
    one = make_waybill(module, ivan, diesel, number="1", date="2026-10-02", season="summer",
                       odometer_out=0, odometer_in=100, fuel_out=50, fuel_in=40)  # fmt: skip
    make_waybill(module, petr, petrol, number="2", date="2026-10-03", season="summer",
                 odometer_out=0, odometer_in=100, fuel_out=10, fuel_in=2)  # fmt: skip
    make_waybill(module, petr, petrol, number="3", date="2026-10-04", season="summer")  # open
    module.data.create("fuel", {"waybill_id": one["id"], "date": "2026-10-02", "liters": 30})


def test_rows_columns_and_totals(module: PrimavtodorModule) -> None:
    fill(module)
    module.settings.set_print_settings(
        {"approver_title": "Начальник", "approver": "А. А. Первый", "composer": "Б. Б. Второй",
         "composer_title": "Инженер", "control": "mechanic"}
    )  # fmt: skip

    content, name = module.fuel_report("2026-10")

    sheet = sheet_of(content)
    assert name == "Анализ расхода ГСМ 2026-10.xlsx"
    assert sheet["F6"].value == "за октябрь 2026 года"
    assert "Начальник" in sheet["I1"].value and "А. А. Первый" in sheet["I1"].value
    by_plate = {sheet.cell(r, 1).value: r for r in range(9, 27) if sheet.cell(r, 1).value}
    assert set(by_plate) == {"Б222ББ125", "Д111ДД125"}  # a vehicle without waybills is left out

    d = by_plate["Д111ДД125"]  # diesel: left column of every pair
    assert (sheet.cell(d, 4).value, sheet.cell(d, 5).value) == ("Иванов И. И.", 100)
    assert [sheet[f"{c}{d}"].value for c in "FHJLN"] == [50, 30, 40, 30, 40]  # 50+30-40 = 40
    assert sheet[f"G{d}"].value is None

    p = by_plate["Б222ББ125"]  # petrol: right column of every pair
    assert [sheet[f"{c}{p}"].value for c in "GIKMO"] == [10, None, 8, 10, 2]
    assert "открытых путевых листов: 1" in sheet[f"R{p}"].value
    assert sheet["A27"].value == "ИТОГО" and sheet["E27"].value == "=SUM(E9:E26)"


def test_more_vehicles_than_lines_move_the_totals_down() -> None:
    rows = [VehicleRow(plate=f"А{i:03d}", model="М", kind="легковой", drivers="Нет водителя",
                       petrol=True, distance_km=1, opening=1.0) for i in range(25)]  # fmt: skip

    sheet = sheet_of(fill_fuel_report(rows, date(2026, 10, 1), ReportSigners(composer="Я")))

    assert sheet["A9"].value == "А000" and sheet["A33"].value == "А024"
    assert sheet["A34"].value == "ИТОГО" and sheet["E34"].value == "=SUM(E9:E33)"
    assert [str(r) for r in sheet.merged_cells.ranges if str(r).startswith("A3")] == ["A34:D34"]
    assert sheet["E36"].value == "Я"


def test_no_data_and_bad_month_are_refused(module: PrimavtodorModule) -> None:
    with pytest.raises(PrintNotAvailableError, match="нет данных"):
        module.fuel_report("2026-10")
    with pytest.raises(PrintNotAvailableError, match="ГГГГ-ММ"):
        module.fuel_report("октябрь")


def test_template_is_blank() -> None:
    sheet = sheet_of(fill_fuel_report([VehicleRow("", "", "легковой", "", True)],
                                      date(2026, 1, 1), ReportSigners()))  # fmt: skip
    texts = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    for word in ("Кихтев", "Матиенко", "Бойко", "LEXUS", "ШМЕЛЬ"):
        assert word.lower() not in texts.lower()


@pytest.fixture
def client(tmp_path: Path):
    app = create_app(Settings(disk_dir=tmp_path / "disk"))
    with TestClient(app, base_url="http://127.0.0.1:8000") as test_client:
        yield test_client


def test_report_over_http(client: TestClient) -> None:
    url = "/api/primavtodor/reports/fuel"
    assert client.get(url, params={"month": "2026-10"}).status_code == 422  # no data yet
    vehicle = client.post("/api/primavtodor/records/vehicles",
                          json={"plate": "А1", "model": "Т"}).json()  # fmt: skip
    driver = client.post("/api/primavtodor/records/employees",
                         json={"full_name": "Иванов Иван", "is_driver": True}).json()  # fmt: skip
    client.post(
        "/api/primavtodor/records/waybills",
        json={"number": "1", "date": "2026-10-07", "driver_id": driver["id"],
              "vehicle_id": vehicle["id"], "odometer_out": 5},
    )  # fmt: skip

    response = client.get(url, params={"month": "2026-10"})

    assert response.status_code == 200 and "filename*=UTF-8''" in response.headers[
        "content-disposition"
    ]
