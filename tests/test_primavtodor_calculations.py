"""Month calculations per vehicle: balances, mileage by the odometer and by waybills, norm."""

from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from tests.test_primavtodor_records import make_driver, make_vehicle, make_waybill
from yukiyasha.config import Settings
from yukiyasha.modules.disk import DiskModule
from yukiyasha.modules.primavtodor import PrimavtodorModule, RecordValidationError
from yukiyasha.web.app import create_app


@pytest.fixture
def module(tmp_path: Path) -> PrimavtodorModule:
    disk = DiskModule(tmp_path / "disk")
    disk.start()
    primavtodor = PrimavtodorModule(disk)
    primavtodor.start()
    return primavtodor


def month_of_two_drivers(module: PrimavtodorModule) -> dict:
    vehicle = make_vehicle(module, plate="А123ВС125", model="TOYOTA", norm_summer="10")
    first = make_driver(module, vehicle["id"], full_name="Иванов Иван Иванович",
                        personnel_number="1", fuel_card_number="1")  # fmt: skip
    second = make_driver(module, vehicle["id"], full_name="Петров Пётр Петрович",
                         personnel_number="2", fuel_card_number="2")  # fmt: skip
    one = make_waybill(module, first, vehicle, number="1", date="2026-10-02", season="summer",
                       odometer_out=1000, odometer_in=1100, fuel_out=20, fuel_in=15)  # fmt: skip
    two = make_waybill(module, first, vehicle, number="2", date="2026-10-05", season="summer",
                       odometer_out=1100, odometer_in=1300, fuel_out=15, fuel_in=10)  # fmt: skip
    make_waybill(module, second, vehicle, number="3", date="2026-10-09", season="summer",
                 odometer_out=1300, odometer_in=1350, fuel_out=10, fuel_in=8)  # fmt: skip
    module.data.create("fuel", {"waybill_id": one["id"], "date": "2026-10-02", "liters": 5})
    module.data.create("fuel", {"waybill_id": two["id"], "date": "2026-10-05", "liters": 10})
    return vehicle


def test_start_end_mileage_actual_and_norm(module: PrimavtodorModule) -> None:
    month_of_two_drivers(module)

    result = module.month_calculations("2026-10")

    (car,) = result["vehicles"]
    assert car["plate"] == "А123ВС125" and car["waybills"] == 3 and car["closed"] == 3
    assert car["first"] == {"number": "1", "date": "2026-10-02"}  # the first day
    assert car["last"] == {"number": "3", "date": "2026-10-09"}  # the last waybill
    assert (car["fuel_start"], car["fills"], car["fuel_end"]) == (20, 15, 8)
    assert (car["odometer_start"], car["odometer_end"], car["km_odometer"]) == (1000, 1350, 350)
    assert car["km_waybills"] == 350
    assert car["consumption"] == 27 == car["consumption_waybills"]  # 20 + 15 - 8
    assert (car["norm"], car["norm_rate"]) == (35, 10)  # 350 km at 10 l per 100 km
    assert car["deviation"] == -8 and car["deviation_pct"] == pytest.approx(-22.9, abs=0.1)
    assert car["notes"] == []
    assert result["totals"]["consumption"] == 27 and result["totals"]["norm"] == 35


def test_per_driver_figures(module: PrimavtodorModule) -> None:
    month_of_two_drivers(module)

    drivers = module.month_calculations("2026-10")["vehicles"][0]["drivers"]

    assert [(d["driver"], d["km"], d["consumption"], d["norm"]) for d in drivers] == [
        ("Иванов Иван Иванович", 300, 25, 30),
        ("Петров Пётр Петрович", 50, 2, 5),
    ]
    assert drivers[0]["deviation"] == -5


def test_unrecorded_mileage_and_broken_remainders_are_named(module: PrimavtodorModule) -> None:
    vehicle = make_vehicle(module, norm_summer="10")
    driver = make_driver(module, vehicle["id"])
    make_waybill(module, driver, vehicle, number="1", date="2026-10-02", season="summer",
                 odometer_out=1000, odometer_in=1100, fuel_out=20, fuel_in=10)  # fmt: skip
    make_waybill(module, driver, vehicle, number="2", date="2026-10-06", season="summer",
                 odometer_out=1200, odometer_in=1300, fuel_out=18, fuel_in=8)  # fmt: skip

    car = module.month_calculations("2026-10")["vehicles"][0]

    assert car["km_odometer"] == 300 and car["km_waybills"] == 200
    assert car["consumption"] == 12 and car["consumption_waybills"] == 20
    notes = " ".join(car["notes"])
    assert "100 км без путевых листов" in notes and "остатки между листами не сходятся" in notes


def test_the_last_closed_waybill_closes_the_month(module: PrimavtodorModule) -> None:
    vehicle = make_vehicle(module, norm_summer="10")
    driver = make_driver(module, vehicle["id"])
    make_waybill(module, driver, vehicle, number="1", date="2026-10-02", season="summer",
                 odometer_out=0, odometer_in=100, fuel_out=20, fuel_in=10)  # fmt: skip
    make_waybill(module, driver, vehicle, number="2", date="2026-10-20", odometer_out=100)  # open

    car = module.month_calculations("2026-10")["vehicles"][0]

    assert car["waybills"] == 2 and car["closed"] == 1
    assert car["last"]["number"] == "1" and car["fuel_end"] == 10
    assert any("Открытых листов: 1" in note for note in car["notes"])


def test_a_vehicle_without_closed_waybills_has_no_end_figures(module: PrimavtodorModule) -> None:
    vehicle = make_vehicle(module)
    driver = make_driver(module, vehicle["id"])
    make_waybill(module, driver, vehicle, number="1", date="2026-10-02")

    car = module.month_calculations("2026-10")["vehicles"][0]

    assert car["last"] is None and car["fuel_end"] is None and car["consumption"] is None
    assert car["deviation"] is None and car["fuel_start"] == 40


def test_mixed_seasons_have_no_single_rate_but_a_summed_norm(module: PrimavtodorModule) -> None:
    vehicle = make_vehicle(module, norm_summer="10", norm_winter="20")
    driver = make_driver(module, vehicle["id"])
    make_waybill(module, driver, vehicle, number="1", date="2026-03-30", season="winter",
                 odometer_out=0, odometer_in=100, fuel_out=30, fuel_in=10)  # fmt: skip
    make_waybill(module, driver, vehicle, number="2", date="2026-03-31", season="summer",
                 odometer_out=100, odometer_in=200, fuel_out=10, fuel_in=0)  # fmt: skip

    car = module.month_calculations("2026-03")["vehicles"][0]

    assert car["norm_rate"] is None and car["norm"] == 30  # 20 + 10 l


def test_empty_month_and_bad_month(module: PrimavtodorModule) -> None:
    assert module.month_calculations("2026-10") == {
        "month": "2026-10",
        "vehicles": [],
        "totals": {"km_waybills": 0, "fuel_start": 0, "fills": 0, "fuel_end": 0,
                   "consumption": 0, "norm": 0, "deviation": 0},
    }  # fmt: skip
    with pytest.raises(RecordValidationError) as exc:
        module.month_calculations("октябрь")
    assert "month" in exc.value.fields


@pytest.fixture
def client(tmp_path: Path):
    app = create_app(Settings(disk_dir=tmp_path / "disk"))
    with TestClient(app, base_url="http://127.0.0.1:8000") as test_client:
        yield test_client


def test_calculations_over_http(client: TestClient) -> None:
    vehicle = client.post("/api/primavtodor/records/vehicles",
                          json={"plate": "А1", "model": "Т", "norm_summer": 10}).json()  # fmt: skip
    driver = client.post("/api/primavtodor/records/employees",
                         json={"full_name": "Иванов Иван", "is_driver": True}).json()  # fmt: skip
    client.post("/api/primavtodor/records/waybills",
                json={"number": "1", "date": "2026-10-05", "driver_id": driver["id"],
                      "vehicle_id": vehicle["id"], "season": "summer", "odometer_out": 0,
                      "odometer_in": 100, "fuel_out": 40, "fuel_in": 30})  # fmt: skip

    ok = client.get("/api/primavtodor/month/2026-10/calculations")
    bad = client.get("/api/primavtodor/month/октябрь/calculations")

    assert ok.status_code == 200 and ok.json()["vehicles"][0]["consumption"] == 10
    assert bad.status_code == 422


# ----- fill-ups by the card of the driver, with or without a waybill -----


def test_fills_by_the_card_count_for_the_vehicle_even_without_a_waybill(
    module: PrimavtodorModule,
) -> None:
    vehicle = make_vehicle(module, plate="Т1", model="Тест", norm_summer="10")
    driver = make_driver(module, vehicle["id"], fuel_card_number="500")
    make_waybill(module, driver, vehicle, number="1", date="2026-10-02", season="summer",
                 odometer_out=0, odometer_in=100, fuel_out=20, fuel_in=15)  # fmt: skip
    one = module.data.create("fuel", {"driver_id": driver["id"], "date": "2026-10-20",
                                      "liters": 30})  # fmt: skip
    assert one["values"]["vehicle_id"] == vehicle["id"] and not one["values"]["waybill_id"]
    assert one["values"]["card_number"] == "500"

    car = module.month_calculations("2026-10")["vehicles"][0]

    assert car["fills"] == 30 and car["fills_without_waybill"] == 1
    assert car["consumption"] == 20 + 30 - 15 == 35  # start + fill-ups - end
    assert any("без путевого листа: 1" in note for note in car["notes"])


def test_a_fill_up_without_a_waybill_needs_a_driver_a_car_and_a_card(
    module: PrimavtodorModule,
) -> None:
    vehicle = make_vehicle(module)
    no_car = make_driver(module, None, fuel_card_number="500", personnel_number="9")
    no_card = make_driver(module, vehicle["id"], fuel_card_number="", personnel_number="8")

    for payload, field in (
        ({"date": "2026-10-20", "liters": 5}, "driver_id"),
        ({"driver_id": no_car["id"], "date": "2026-10-20", "liters": 5}, "vehicle_id"),
        ({"driver_id": no_card["id"], "date": "2026-10-20", "liters": 5}, "driver_id"),
    ):
        with pytest.raises(RecordValidationError) as exc:
            module.data.create("fuel", payload)
        assert field in exc.value.fields
    given = module.data.create("fuel", {"driver_id": no_car["id"], "vehicle_id": vehicle["id"],
                                        "date": "2026-10-20", "liters": 5})  # fmt: skip
    assert given["values"]["vehicle_id"] == vehicle["id"]


def test_a_car_with_fill_ups_but_no_waybills_is_listed_and_carries_its_balance(
    module: PrimavtodorModule,
) -> None:
    vehicle = make_vehicle(module, plate="Т2", model="Тест")
    driver = make_driver(module, vehicle["id"], fuel_card_number="501")
    make_waybill(module, driver, vehicle, number="1", date="2026-09-20", season="summer",
                 odometer_out=0, odometer_in=100, fuel_out=20, fuel_in=12)  # September
    module.data.create("fuel", {"driver_id": driver["id"], "date": "2026-10-05", "liters": 40})

    car = module.month_calculations("2026-10")["vehicles"][0]

    assert car["waybills"] == 0 and car["first"] is None and car["last"] is None
    assert car["fuel_start"] == 12 and car["fills"] == 40  # the remainder of the last waybill
    assert car["consumption"] is None
    assert any("нет путевых листов" in note for note in car["notes"])
    findings = module.month_review("2026-10")["findings"]
    assert any(f["code"] == "FUEL_NO_WAYBILLS" for f in findings)


def test_card_sheet_and_analysis_use_the_same_fill_ups(module: PrimavtodorModule) -> None:
    import io

    import openpyxl

    vehicle = make_vehicle(module, plate="Т3", model="Тест", norm_summer="10")
    driver = make_driver(module, vehicle["id"], fuel_card_number="502")
    make_waybill(module, driver, vehicle, number="1", date="2026-10-02", season="summer",
                 odometer_out=0, odometer_in=100, fuel_out=20, fuel_in=15)  # fmt: skip
    module.data.create("fuel", {"driver_id": driver["id"], "date": "2026-10-20", "liters": 30})

    card = openpyxl.load_workbook(io.BytesIO(module.fuel_cards(vehicle["id"], "2026-10")[0]))
    report = openpyxl.load_workbook(io.BytesIO(module.fuel_report("2026-10")[0]))["Анализ"]

    assert [card.worksheets[0].cell(r, 2).value for r in (9, 10)] == [30, None]
    petrol_or_diesel = [report[f"{c}9"].value for c in "HI"]  # the fill-ups column pair
    assert 30 in petrol_or_diesel
