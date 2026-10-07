"""Timesheet on the form Т-12: marks, hours, painted days off, more than one page of people."""

import io
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from tests.test_primavtodor_records import make_driver, make_vehicle, make_waybill
from yukiyasha.config import Settings
from yukiyasha.modules.disk import DiskModule
from yukiyasha.modules.primavtodor import PrimavtodorModule, PrintNotAvailableError
from yukiyasha.modules.primavtodor.timesheet_form import (
    BLOCKS,
    TimesheetForm,
    TimesheetPerson,
    _column,
    fill_timesheet,
)
from yukiyasha.web.app import create_app

SETTINGS = {"org_name": 'АО "ТЕСТ"', "unit_name": "Гараж", "composer_title": "Инженер",
            "composer": "Я. Я.", "approver_title": "Начальник", "approver": "Н. Н.",
            "control": "mechanic"}  # fmt: skip


@pytest.fixture
def module(tmp_path: Path) -> PrimavtodorModule:
    disk = DiskModule(tmp_path / "disk")
    disk.start()
    primavtodor = PrimavtodorModule(disk)
    primavtodor.start()
    return primavtodor


def sheet_of(content: bytes):
    return openpyxl.load_workbook(io.BytesIO(content))["Табель"]


def cell(sheet, row: int, day: int):
    return sheet.cell(row, _column(day))


def test_marks_hours_and_painted_days_off(module: PrimavtodorModule) -> None:
    module.settings.set_print_settings(SETTINGS)
    vehicle = make_vehicle(module)
    driver = make_driver(module, vehicle["id"], full_name="Иванов Иван Иванович",
                         personnel_number="77")  # fmt: skip
    make_waybill(module, driver, vehicle, number="1", date="2026-10-07")  # Wednesday
    make_waybill(module, driver, vehicle, number="2", date="2026-10-10", odometer_out=1)  # Sat
    module.timesheet.set_mark("2026-10", driver["id"], "2026-10-12", "ОТ")

    content, name = module.timesheet_form("2026-10")

    sheet = sheet_of(content)
    assert name == "Табель 2026-10.xlsx"
    assert sheet["I2"].value == 'АО "ТЕСТ"' and sheet["I4"].value == "Гараж"
    assert sheet["F16"].value == "Иванов И. И. Водитель" and sheet["V16"].value == "77"
    assert [cell(sheet, 16, d).value for d in (7, 10, 12, 8)] == ["Я", "РВ", "ОТ", None]
    assert [cell(sheet, 17, d).value for d in (7, 10, 12)] == [8, 8, None]
    assert cell(sheet, 12, 10).fill.fgColor.rgb == "FFFFFF00"  # Saturday header is painted
    assert cell(sheet, 16, 10).fill.fgColor.rgb == "FFFFCC00"  # worked on a day off
    assert cell(sheet, 16, 7).fill.fill_type is None  # an ordinary working day
    assert sheet["U48"].value == "Инженер" and sheet["BW48"].value == "Я. Я."
    assert sheet["GJ48"].value == "Н. Н." and sheet["IB48"].value.isdigit()
    assert sheet["CO17"].value == "=SUM(AC17:CN17)"  # the form's own totals stay live


def test_the_shortened_day_counts_seven_hours(module: PrimavtodorModule) -> None:
    vehicle = make_vehicle(module)
    driver = make_driver(module, vehicle["id"])
    make_waybill(module, driver, vehicle, number="1", date="2026-04-30")  # before 1 May

    sheet = sheet_of(module.timesheet_form("2026-04")[0])

    assert cell(sheet, 16, 30).value == "Я" and cell(sheet, 17, 30).value == 7


def test_a_holiday_that_is_a_weekday_is_painted(module: PrimavtodorModule) -> None:
    vehicle = make_vehicle(module)
    make_driver(module, vehicle["id"])

    sheet = sheet_of(module.timesheet_form("2026-05")[0])

    assert cell(sheet, 12, 11).fill.fgColor.rgb == "FFFFFF00"  # Monday 11 May, transfer
    assert cell(sheet, 12, 12).fill.fill_type is None


def test_more_people_than_the_form_holds_repeat_the_block() -> None:
    people = [TimesheetPerson(name=f"Человек {i}", marks={1: "Я"}, hours={1: 8})
              for i in range(BLOCKS + 3)]  # fmt: skip
    form = TimesheetForm(month=__import__("datetime").date(2026, 10, 1), days_in_month=31,
                         off_days={3, 4}, persons=people, composer="Составитель")  # fmt: skip

    sheet = sheet_of(fill_timesheet(form))

    last = 16 + 2 * (BLOCKS + 2)
    assert sheet[f"F{last}"].value == "Человек 17" and cell(sheet, last, 1).value == "Я"
    assert cell(sheet, last + 1, 1).value == 8
    assert sheet[f"A{last}"].value == BLOCKS + 3
    assert sheet[f"BW{48 + 6}"].value == "Составитель"  # the signature moved down by 3 blocks
    assert sheet["CO" + str(last + 1)].value == f"=SUM(AC{last + 1}:CN{last + 1})"


def test_the_blank_form_is_clean_and_empty_lines_show_no_totals(
    module: PrimavtodorModule,
) -> None:
    vehicle = make_vehicle(module)
    make_driver(module, vehicle["id"], full_name="Единственный Сотрудник Сотрудникович")
    sheet = sheet_of(module.timesheet_form("2026-10")[0])

    texts = " ".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
    for word in ("Кихтев", "Матиенко", "Мальцев", "Бойко", "Жариков"):
        assert word not in texts
    assert sheet["CO18"].value is None and sheet["FV18"].value is None  # an empty line


def test_no_employees_is_refused(module: PrimavtodorModule) -> None:
    with pytest.raises(PrintNotAvailableError, match="нет сотрудников"):
        module.timesheet_form("2026-10")


@pytest.fixture
def client(tmp_path: Path):
    app = create_app(Settings(disk_dir=tmp_path / "disk"))
    with TestClient(app, base_url="http://127.0.0.1:8000") as test_client:
        yield test_client


def test_timesheet_form_over_http(client: TestClient) -> None:
    client.post("/api/primavtodor/records/employees",
                json={"full_name": "Иванов Иван", "is_driver": True})  # fmt: skip

    ok = client.get("/api/primavtodor/timesheet/form", params={"month": "2026-10"})
    bad = client.get("/api/primavtodor/timesheet/form", params={"month": "октябрь"})

    assert ok.status_code == 200 and "filename*=UTF-8''" in ok.headers["content-disposition"]
    assert ok.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert bad.status_code == 422
